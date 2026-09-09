import unittest
from pathlib import Path

from openpyxl import load_workbook

import run_audit


class ReportPresentationTests(unittest.TestCase):
    def make_result(self):
        issue = lambda priority, rule_id, row, code, message: {
            "status": "open",
            "priority": priority,
            "rule_id": rule_id,
            "excel_row": row,
            "kks_code": code,
            "message": message,
            "suggestion": "请核对后整改。",
            "evidence": "测试证据",
            "final_decision": "confirmed_issue",
        }
        return {
            "source_file": "sample.xlsx",
            "sheet": "Sheet1",
            "header_row": 1,
            "data_rows": 2,
            "columns": ["KKS", "父级", "名称"],
            "sheet_info": [{"sheet": "Sheet1", "max_row": 3, "max_col": 3, "nonempty_rows": 2, "first_nonempty_row": 1, "last_nonempty_row": 3, "header_preview": ["KKS", "父级", "名称"]}],
            "metrics": {"unique_codes": 2, "duplicate_code_groups": 0, "orphan_rows": 0, "prefix_mismatch_rows": 0, "long_code_rows": 1},
            "priority_counts": {"P0": 1, "P1": 1, "P2": 1},
            "issue_count": 3,
            "resolved_count": 0,
            "conclusion": "测试结论",
            "structural_notes": [],
            "issues": [
                issue("P0", "KKS-04b", 2, "01AAA10QM0O1", "字符错误"),
                issue("P1", "KKS-17/22", 3, "01AAA10QM001A", "扩展编码"),
                issue("P2", "KKS-18", 4, "01AAA10QM002", "历史迁移"),
            ],
            "resolved_reviews": [],
            "import_governance": {"direct_import_ready": False},
            "scope_comparison": None,
            "_ai_context": {"records": [{"kks_code": "01AAA10QM001"}, {"kks_code": "01AAA10QM001A"}]},
        }

    def test_business_metrics_use_real_priority_counts(self):
        result = self.make_result()
        self.assertEqual(
            run_audit.quality_metrics(result),
            {
                "effective_kks": 2,
                "device_level_codes": 1,
                "duplicate_codes": 0,
                "parent_errors": 0,
                "p0": 1,
                "p1": 1,
                "p2": 1,
                "rule_p1": 1,
                "rule_p2": 1,
                "focused_p1": 1,
                "focused_p2": 1,
            },
        )

    def test_problem_workbook_uses_business_group_sheets(self):
        path = Path.cwd() / "outputs" / "test_report_presentation.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            run_audit.write_issue_workbook_xlsx(path, self.make_result())
            workbook = load_workbook(path, read_only=False, data_only=True)
            # 广元风格 sheet 结构：概览 + 按业务分组的明细 sheet + 结构说明
            self.assertIn("概览", workbook.sheetnames)
            self.assertIn("P0_阻断", workbook.sheetnames)
            self.assertIn("P2_其他提示", workbook.sheetnames)
            self.assertIn("结构说明与范围", workbook.sheetnames)
            self.assertTrue(workbook["概览"]["A1"].value.startswith("sample 审核总览"))
            # P0_阻断 sheet：第 3 行为表头、数据从第 4 行开始（与广元样式一致）
            headers = [cell.value for cell in workbook["P0_阻断"][3]]
            self.assertEqual(headers, ["行", "编码", "规则", "问题", "整改建议"])
            first = [cell.value for cell in workbook["P0_阻断"][4]]
            self.assertEqual(first, [2, "01AAA10QM0O1", "KKS-04b", "字符错误", "请核对后整改。"])
        finally:
            path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
