# -*- coding: utf-8 -*-
"""交付列「KKS编码（修编后）」**全表唯一** —— 端到端回归（国标 3.1.1）。

## 口径（用户原话）

> 一条重复的 kks 编码都不能有 你不知道 kks 编码是唯一的吗？

KKS 编码是**全厂唯一标识**（GB/T 50549-2020 3.1.1），交付列里**一条重复都不许有** ——
包括**源表自带**的重复：两个不同设备被给了同一个原KKS编码时，交付码必然相撞，
不能因为"不是引擎造的"就放过。

判据落在 `fill_original_table` 的**写出级撞码闸门**上（判实际写出去的值），
因为引擎内部那道唯一性校验早于结构闸门与「指令/测点不写」两条规则，覆盖不到。
"""
from __future__ import annotations

import contextlib
import io
import pathlib
import sys
import tempfile
import unittest
from collections import Counter

import openpyxl

PROJECT = pathlib.Path(__file__).resolve().parents[1]
for p in (str(PROJECT), str(PROJECT / "tools"), str(PROJECT / "kks_gen")):
    if p not in sys.path:
        sys.path.insert(0, p)

import gen_step1  # noqa: E402

HEADER = ["序号", "机组号", "系统", "设备类型", "原设备名称",
          "设备名称（按标准修改加机组号）*", "原KKS编码", "KKS编码（修编后）*",
          "父级KKS编码*", "修改意见（审核人填写）"]

# 🔴 源表的真实现象：**同一个原KKS编码给了两个不同设备**（实测行 43024 / 56536）
SHARED = "01ABC01WA001"
ROWS = [
    ("省煤器出口压力变送器1", "01HAC13BP101"),
    ("#1机组直流配电设备A", SHARED),
    ("#1机组直流配电设备B", SHARED),
]


class DeliveredCodeUniqueEndToEndTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        root = pathlib.Path(cls._tmp.name)
        cls.src = root / "提资.xlsx"
        cls.out_dir = root / "out"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "sheet1"
        ws.append(["一、表格说明：该表是依据提资内容整理的专业设备汇总表"])
        ws.append(HEADER)
        for i, (name, code) in enumerate(ROWS, start=1):
            ws.append([i, "1号机组", "", "", name, "", code, "", "", ""])
        wb.save(cls.src)
        wb.close()
        cls.original_bytes = cls.src.read_bytes()

        argv = sys.argv
        sys.argv = ["gen_step1.py", str(cls.src), str(cls.out_dir), "--no-ai"]
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                gen_step1.main()
        finally:
            sys.argv = argv
        cls.rows = cls._read(cls.out_dir / f"{cls.src.stem}_已填写.xlsx")

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    @staticmethod
    def _read(path: pathlib.Path) -> list[dict]:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        head = next(r for r in ws.iter_rows(min_row=1, max_row=6, values_only=True)
                    if r and "原KKS编码" in [str(c or "").strip() for c in r])
        idx = {str(h).strip(): i for i, h in enumerate(head) if str(h or "").strip()}
        out = [{h: (r[i] if i < len(r) else None) for h, i in idx.items()}
               for r in ws.iter_rows(min_row=3, values_only=True)]
        wb.close()
        return [r for r in out if r.get("原KKS编码")]

    def _codes(self) -> list[str]:
        return [str(r.get("KKS编码（修编后）*") or "").strip() for r in self.rows
                if str(r.get("KKS编码（修编后）*") or "").strip()]

    def test_source_file_untouched(self):
        self.assertEqual(self.original_bytes, self.src.read_bytes())

    def test_no_duplicate_code_anywhere(self):
        """🔴 契约本体：交付列**一条重复都没有**（源表自带的重复也要拦住）。"""
        dup = {c: n for c, n in Counter(self._codes()).items() if n > 1}
        self.assertEqual({}, dup, f"交付编码列出现重复值：{dup}")

    def test_colliding_rows_are_both_left_blank(self):
        """撞码的行**一律留空** —— 不替引擎挑一个留下（它无从判断该给谁）。"""
        sharing = [r for r in self.rows if str(r.get("原KKS编码")).strip() == SHARED]
        self.assertEqual(2, len(sharing), "夹具里应有两条同原码的行")
        for r in sharing:
            with self.subTest(name=str(r.get("原设备名称"))[:20]):
                self.assertIn(r.get("KKS编码（修编后）*"), (None, ""),
                              "同原码的两行都不能拿到交付码")

    def test_collision_is_explained_not_silent(self):
        """留空必须写明原因（含国标 3.1.1 与撞哪一行），不能静默留白。"""
        for r in [x for x in self.rows if str(x.get("原KKS编码")).strip() == SHARED]:
            note = str(r.get("修编说明") or "")
            with self.subTest(note=note[:40]):
                self.assertIn("3.1.1", note)
                self.assertIn("全厂范围内唯一", note)

    def test_unique_row_keeps_its_code(self):
        """⚠️ 不能把契约执行成"整列都不填"：不撞码的行照旧要有码。"""
        r = next(x for x in self.rows if str(x.get("原KKS编码")).strip() == "01HAC13BP101")
        self.assertEqual("01HAC13BP101", str(r.get("KKS编码（修编后）*") or "").strip())


if __name__ == "__main__":
    unittest.main()
