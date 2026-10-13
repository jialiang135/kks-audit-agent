"""Regression coverage for the shared compliance gate and actual delivery."""
import concurrent.futures
import pathlib
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
for directory in (ROOT, ROOT / "tools", ROOT / "kks_gen"):
    sys.path.insert(0, str(directory))
import gen_step1
from kks_gen.compliance import assess, unit_conflicts
from kks_gen.gb_rules import unit_label_to_plant_code
from revise import revise_code
from tree_rules import build_tree


class CompliancePipelineTests(unittest.TestCase):
    def test_optional_plant_characters(self):
        for code in ("01HAL11GP001", "K01HAL11GP001", "011HAL11GP001", "K011HAL11GP001"):
            with self.subTest(code=code):
                self.assertTrue(assess(code).writable)

    def test_strict_raw_characters_and_suffix(self):
        for code in (" 01HAL11GP001", "01hal11gp001", "01HAL11GP001.A", "01HAL11GP000"):
            with self.subTest(code=code):
                self.assertFalse(assess(code).writable)

    def test_all_product_segments_checked(self):
        self.assertTrue(assess("01HAL11GP001-MA01").writable)
        self.assertFalse(assess("01HAL11GP001-MA01-IO01").writable)

    def test_terminal_io_exception(self):
        self.assertTrue(assess(":IO").writable)
        self.assertFalse(assess("01HAL11QI001").writable)

    def test_shared_unit_requires_agreement(self):
        self.assertEqual("", unit_label_to_plant_code("1、2、3号机组共用"))
        self.assertTrue(assess("91HAL11GP001").pending)
        self.assertTrue(revise_code("91HAL11GP001").needs_review)

    def test_unit_conflict_never_reassigns(self):
        result = revise_code("01HAL11GP001", unit_label="2号机组")
        self.assertEqual("unresolved", result.status)
        self.assertEqual("", result.revised)
        self.assertTrue(unit_conflicts("K01HAL11GP001", "二号机组给水泵"))

    def test_parallel_arguments_and_callbacks_are_isolated(self):
        barrier = threading.Barrier(2)
        events = {"first": [], "second": []}
        original_argv = list(sys.argv)
        def fake_main(argv):
            barrier.wait(timeout=5)
            gen_step1._emit({"job": argv[0]})
            return 0
        with patch.object(gen_step1, "_main", fake_main):
            with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
                futures = [pool.submit(gen_step1.main, [job], progress_callback=events[job].append)
                           for job in events]
                self.assertEqual([0, 0], [f.result() for f in futures])
        self.assertEqual(original_argv, sys.argv)
        self.assertEqual([{"job": "first"}], events["first"])
        self.assertEqual([{"job": "second"}], events["second"])

    def test_ai_cannot_dismiss_proven_structural_error(self):
        import run_audit
        issue = {"rule_id": "KKS-04b", "compliance_locked": True, "status": "open"}
        result = {"issues": [issue], "conclusion": "规则审核"}
        def fake_review(result, **kwargs):
            issue["final_decision"] = "likely_false_positive"
            issue["status"] = "resolved"
            return {"status": "completed"}
        with patch.object(run_audit, "review_issue_candidates", fake_review), \
             patch.object(run_audit, "summarize_audit", return_value={}):
            run_audit.apply_ai_review_stage(result)
        self.assertEqual("confirmed_issue", issue["final_decision"])
        self.assertEqual("open", issue["status"])

    def test_audit_suggestions_never_contain_replacement_codes(self):
        import run_audit
        issue = {"rule_id": "KKS-42", "suggestion": "核对01HAL11GP001及61LBG82",
                 "kks_code": "01HAL11GP001", "evidence_missing": True}
        result = {"issues": [issue], "conclusion": "规则审核"}
        with patch.object(run_audit, "review_issue_candidates", return_value={
                "status": "completed", "results": [{"suggestion": "改为61LBG82"}]}), \
             patch.object(run_audit, "summarize_audit", return_value={}):
            run_audit.apply_ai_review_stage(result)
        self.assertNotIn("01HAL11GP001", issue["suggestion"])
        self.assertNotIn("61LBG82", issue["suggestion"])
        self.assertEqual("01HAL11GP001", issue["kks_code"])
        self.assertEqual("needs_human", issue["final_decision"])

    def test_action_target_is_not_equipment_subject(self):
        from kks_gen.audit_policy import subject
        for name in ("炉膛压力高跳送风机", "炉膛压力低跳引风机", "联锁停止给水泵"):
            self.assertIsNone(subject(name))
        self.assertIsNotNone(subject("A送风机"))

    def test_batch_evidence_pool_preserves_every_row(self):
        import ai_review
        shared = {"sheet": "设备", "excel_row": 2, "kks_code": "01HAL11"}
        current = {"sheet": "设备", "excel_row": 3, "kks_code": "01HAL11GP001"}
        group = [{"issue_index": 0, "evidence_packet": {"current_row": current,
                  "parent_rows": [shared], "same_level_rows": [current], "adjacent_rows": [shared]}}]
        request = ai_review._group_evidence_request(group)
        packet = request["candidates"][0]["evidence_packet"]
        self.assertEqual(2, len(request["row_pool"]))
        self.assertEqual(current, request["row_pool"][packet["current_row"]])
        self.assertEqual(shared, request["row_pool"][packet["parent_rows"][0]])
        self.assertEqual([shared], group[0]["evidence_packet"]["parent_rows"])

    def test_incomplete_ai_batch_is_reported_and_never_passes(self):
        import ai_review
        with patch.object(ai_review, "_review_group", return_value=[]):
            reviews, errors = ai_review._review_groups_parallel(None, "test", [[{"issue_index": 0}]],
                verification=False, audit_profile={}, global_review={}, progress_callback=None,
                candidate_count=1, reviewed_count=0, percent_start=0, percent_end=100)
        self.assertFalse(reviews)
        self.assertTrue(errors)
        final = ai_review._default_final_review({"status": "needs_review"}, candidate=True)
        self.assertEqual("needs_human", final["final_decision"])

    def test_finished_columns_exclude_original_and_auxiliary_fields(self):
        import run_audit
        header = ("KKS编码", "原KKS编码", "原设备名称", "机组号", "设备类型", "系统",
                  "新KKS编码（修编后）*", "新设备名称（修改后）", "父级KKS编码*")
        columns = run_audit.audit_columns(header)
        row = ("WRONG", "BROKEN", "原名称", "2号机", "错误类别", "错误系统",
               "01HAL11GP001", "给水泵A", "01HAL11")
        rec = run_audit.make_record(2, row, columns)
        self.assertEqual("01HAL11GP001", rec["kks_code"])
        self.assertEqual("给水泵A", rec["name"])
        for key in ("old_code", "unit", "device_type", "system"):
            self.assertEqual("", rec[key])
        self.assertNotIn("原KKS编码", run_audit._template_columns(header, columns))

    def test_finished_blank_name_never_falls_back_to_original(self):
        import run_audit
        header = ("原设备名称", "新KKS编码", "新设备名称", "父级KKS编码")
        rec = run_audit.make_record(2, ("原来有名称", "01HAL11GP001", "", "01HAL11"), run_audit.audit_columns(header))
        self.assertEqual("", rec["name"])

    def test_finished_quantity_name_does_not_prove_signal_role(self):
        import run_audit
        rec = run_audit.make_record(2, ("01HAL11BP001", "出口压力", "01HAL11"),
            run_audit.audit_columns(("新KKS编码", "新设备名称", "父级KKS编码")))
        self.assertFalse(run_audit.record_role(rec).is_point)
        rec["name"] = "出口压力反馈"
        self.assertTrue(run_audit.record_role(rec).is_point)

    def test_report_artifacts_do_not_exist_before_ai_finishes(self):
        import run_audit
        with tempfile.TemporaryDirectory() as folder:
            path = pathlib.Path(folder) / "finished.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.append(["新KKS编码", "新设备名称", "父级KKS编码"])
            workbook.active.append(["01HAL11GP001", "给水泵A", "01HAL11"])
            workbook.save(path)
            workbook.close()
            events = []
            def check_before_review(result, **kwargs):
                html_name, excel_name = run_audit.output_artifact_names(path)
                self.assertFalse((path.parent / html_name).exists())
                self.assertFalse((path.parent / excel_name).exists())
                return {"status": "completed", "candidate_count": 0, "reviewed_count": 0}
            with patch.object(run_audit, "review_issue_candidates", side_effect=check_before_review), \
                 patch.object(run_audit, "summarize_audit", return_value={}), \
                 patch.object(run_audit, "record_and_render", return_value=path):
                result = run_audit.audit_file(path, path.parent, run_ai_review=True,
                    write_outputs=True, progress_callback=events.append)
            self.assertFalse(any(event.get("stage") == "rule_report_ready" for event in events))
            self.assertNotIn("stage_scope_note", result)

    def test_measurement_qualifier_preserves_physical_valve(self):
        from kks_gen.row_policy import classify_row
        name = "1号机除氧器远传模拟量液位1汽侧一次门"
        for tail in ("", "2", "A"):
            self.assertFalse(classify_row("01LAA10QM411", name + tail).is_point)
        for tail in ("已关闭", "关闭反馈", "已开启", "开到位", "关闭", "关闭指令", "运行", "故障", "允许"):
            self.assertTrue(classify_row("01LAA10QM411", name + tail).is_point)
        self.assertTrue(classify_row("01LAA10QM411", "1号机关闭一次门").is_point)
        self.assertFalse(classify_row("01HAL11BP001", "模拟量压力变送器").is_point)
        self.assertTrue(classify_row("01HAL11BP001", "模拟量压力变送器故障").is_point)

    def test_ai_evidence_uses_excel_rows_not_pool_identifiers(self):
        import ai_review
        from unittest.mock import Mock
        client = Mock()
        client.review.return_value = {"reviews": [{"issue_index": 0, "decision": "needs_human",
            "evidence": "当前r0，父r1", "reason": "核对r0", "suggestion": "核对资料"}]}
        reviews = ai_review._review_group(client, "test", [{"issue_index": 0, "evidence_packet": {
            "current_row": {"sheet": "汇总", "excel_row": 76},
            "parent_rows": [{"sheet": "汇总", "excel_row": 60}]}}])
        self.assertEqual("当前汇总!76，父汇总!60", reviews[0]["evidence"])

    def test_solenoid_valve_qa_is_supported_by_standard_examples(self):
        from kks_gen.audit_policy import findings, subject
        import ai_review
        self.assertIn("QA", subject("2号机油压试验装置电磁阀")[1])
        rec = {"kks_code": "02MAK82QA403", "name": "2号机油压试验装置电磁阀", "parent_code": "02MAK82", "finished_product": True}
        self.assertFalse(any(stage == 1 and confirmed for stage, _, _, confirmed in findings(rec)))
        self.assertIn("电磁阀", ai_review._classification_reference(rec["kks_code"])["device"]["application_examples"])

    def test_formal_report_requires_complete_ai_review(self):
        from agent_runtime.tools import ReportTool
        for status, count, reviewed in (("disabled", 0, 0), ("partial", 2, 1), ("completed", 2, 1)):
            with self.subTest(status=status), self.assertRaises(RuntimeError):
                ReportTool().execute(pathlib.Path("input.xlsx"), pathlib.Path("outputs"),
                    {"ai_review": {"status": status, "candidate_count": count, "reviewed_count": reviewed}})

    def test_completed_audit_status_survives_service_restart(self):
        import server
        from types import SimpleNamespace
        with tempfile.TemporaryDirectory() as folder:
            fake = SimpleNamespace(jobs={}, jobs_lock=threading.Lock(), runs_dir=pathlib.Path(folder))
            run_id = "20261013T085854_9cd2838d"
            server._update_job_state(fake, run_id, status="completed", ai_review={"status": "partial"})
            fake.jobs.clear()
            state = server._get_job_state(fake, run_id)
            self.assertEqual("completed", state["status"])
            self.assertEqual("partial", state["ai_review"]["status"])

    def test_model_concurrency_limit_is_shared_by_clients(self):
        import ai_review, time
        activity = {"active": 0, "peak": 0}
        lock = threading.Lock()
        class Response:
            def __enter__(self):
                with lock:
                    activity["active"] += 1
                    activity["peak"] = max(activity["peak"], activity["active"])
                return self
            def read(self):
                time.sleep(0.02)
                return b'{}'
            def __exit__(self, *args):
                with lock:
                    activity["active"] -= 1
        config = ai_review.AIConfig("https://example.invalid", "test", "test", 10, True)
        clients = [ai_review.OpenAICompatibleClient(config) for _ in range(2)]
        with patch.object(ai_review, "MODEL_REQUEST_SLOTS", threading.BoundedSemaphore(2)), \
             patch.object(ai_review.urllib.request, "urlopen", side_effect=lambda *a, **k: Response()):
            with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
                results = list(pool.map(lambda i: clients[i % 2]._request_json("test", {}), range(8)))
        self.assertEqual([{}] * 8, results)
        self.assertLessEqual(activity["peak"], 2)

    def test_transport_error_does_not_trigger_second_retry_layer(self):
        import ai_review
        from unittest.mock import Mock
        client = Mock()
        client.review.side_effect = ai_review.AIReviewError("模型接口请求超时")
        with self.assertRaises(ai_review.AIReviewError):
            ai_review._review_group(client, "test", [{"issue_index": 0}])
        self.assertEqual(1, client.review.call_count)

    def test_missing_classification_evidence_skips_ai_and_remains_pending(self):
        import ai_review, run_audit
        with tempfile.TemporaryDirectory() as folder:
            path = pathlib.Path(folder) / "audit.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["KKS编码", "设备名称", "设备类型"])
            sheet.append(["01HAL11GP001", "待确认设备", "综合热工设备"])
            workbook.save(path)
            workbook.close()
            result = run_audit.audit_file(path, pathlib.Path(folder), run_ai_review=False, write_outputs=False)
            findings = [i for i in result["issues"] if i["rule_id"] in {"KKS-40", "KKS-41", "KKS-42"}]
            self.assertEqual(3, len(findings))
            for finding in findings:
                self.assertTrue(finding["evidence_missing"])
                self.assertFalse(ai_review._is_candidate(finding))
                self.assertEqual("needs_human", finding["final_decision"])

    def _deliver(self, folder, *, existing=False, code="01HAL11GP001", unresolved=False):
        source = pathlib.Path(folder) / "input.xlsx"
        output = pathlib.Path(folder) / "output.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["完整保留此说明", None, None, None, None])
        sheet.merge_cells("A1:E1")
        sheet.append(["原KKS编码", "原设备名称", "KKS编码（修编后）*", "设备名称（按标准修改加机组号）*", "父级KKS编码*"])
        for name in ("给水泵A", "给水泵B"):
            sheet.append([code, name, code if existing else "", "", ""])
        workbook.save(source)
        workbook.close()
        original = source.read_bytes()
        seeds, layout = gen_step1.load(source)
        if unresolved:
            for seed in seeds:
                seed["revise_status"] = "unresolved"
                seed["revise_note"] = "机组归属冲突"
        nodes = build_tree(seeds)["nodes"]
        result = gen_step1.fill_original_table(source, output, nodes, layout, seeds=seeds)
        self.assertEqual(original, source.read_bytes())
        return output, result

    def test_unconfirmed_or_unresolved_code_is_not_newly_filled(self):
        for code, unresolved in (("91HAL11GP001", False), ("01HAL11GP001", True), ("01BUC16WA001", False)):
            with self.subTest(code=code, unresolved=unresolved), tempfile.TemporaryDirectory() as folder:
                output, result = self._deliver(folder, code=code, unresolved=unresolved)
                workbook = openpyxl.load_workbook(output)
                original_rows = [r for r in workbook.active.iter_rows(min_row=3, values_only=True) if r[0]]
                self.assertTrue(all(r[3] is None for r in original_rows))
                self.assertEqual("pending", result["compliance"]["status"])
                workbook.close()

    def test_merged_description_and_empty_collision(self):
        with tempfile.TemporaryDirectory() as folder:
            output, result = self._deliver(folder)
            workbook = openpyxl.load_workbook(output)
            sheet = workbook.active
            self.assertEqual("完整保留此说明", sheet["B1"].value)
            self.assertIn("B1:F1", [str(r) for r in sheet.merged_cells.ranges])
            original_rows = [r for r in sheet.iter_rows(min_row=3, values_only=True) if r[0]]
            self.assertTrue(all(r[3] is None for r in original_rows))
            self.assertEqual(2, result["dup_blocked_rows"])
            workbook.close()

    def test_preexisting_duplicates_are_cleared(self):
        with tempfile.TemporaryDirectory() as folder:
            output, result = self._deliver(folder, existing=True)
            self.assertEqual("pending", result["compliance"]["status"])
            workbook = openpyxl.load_workbook(output)
            original_rows = [r for r in workbook.active.iter_rows(min_row=3, values_only=True) if r[0]]
            self.assertTrue(all(r[3] is None for r in original_rows))
            workbook.close()


if __name__ == "__main__":
    unittest.main()
