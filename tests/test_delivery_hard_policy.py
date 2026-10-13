"""Latest user contract: no point codes, and no final-code duplicates."""
import contextlib
import io
import pathlib
import sys
import tempfile
import unittest
from collections import Counter

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
import gen_step1
import run_audit
from kks_gen.row_policy import classify_row

HEADER = ["原KKS编码", "原设备名称", "设备类型", "KKS编码（修编后）*", "设备名称", "父级KKS编码*"]


class DeliveryHardPolicyTests(unittest.TestCase):
    def test_row_roles_include_optional_plant_prefix_and_physical_exception(self):
        for code, name, dtype in (("K011HAL11GP001XB21", "运行反馈", ""),
                                  ("01HAL11BP101", "出口压力1", ""),
                                  ("", "启动给水泵", ""), ("01HAL11GP001", "控制状态", "DI")):
            self.assertTrue(classify_row(code, name, dtype).is_point)
        self.assertFalse(classify_row("01HAL11BP101", "压力变送器1").is_point)

    def test_prefilled_point_codes_and_all_duplicate_participants_are_cleared(self):
        with tempfile.TemporaryDirectory() as directory:
            source = pathlib.Path(directory) / "source.xlsx"
            folder = pathlib.Path(directory) / "out"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(HEADER)
            rows = [
                ["01HAL11GP001", "给水泵A", "泵", "01HAL11GP001", "给水泵A", "01HAL11"],
                ["01HAL11GP001XB21", "给水泵A已运行", "DI", "01HAL11GP001", "运行反馈", "01HAL11"],
                ["K011HAL11GP001YB21", "启动给水泵", "DO", "K011HAL11GP001", "启动指令", "K011HAL11"],
                ["01HAL11BP101", "出口压力1", "测点", "01HAL11BP101", "压力1", "01HAL11"],
                ["01HAL11BP102", "压力变送器2", "仪表", "01HAL11BP102", "压力变送器2", "01HAL11"],
                ["01HAL11GP002", "给水泵B", "泵", "01HAL11GP002", "给水泵B", "01HAL11"],
                ["01HAL11GP002", "给水泵C", "泵", "01HAL11GP002", "给水泵C", "01HAL11"],
                ["", "频率反馈", "AI", "", "", "01HAL11"],
            ]
            for row in rows:
                sheet.append(row)
            other = workbook.create_sheet("第二专业")
            other.append(HEADER)
            other.append(["01HAL11GP001", "另一设备", "泵", "01HAL11GP001", "另一设备", "01HAL11"])
            workbook.save(source)
            workbook.close()
            before = source.read_bytes()
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(0, gen_step1.main([str(source), str(folder), "--no-ai"]))
            self.assertEqual(before, source.read_bytes())
            delivered = openpyxl.load_workbook(folder / "source_已填写.xlsx")
            primary = delivered.worksheets[0]
            by_original_row = {row[0]: row for row in primary.iter_rows(min_row=2, values_only=True) if row[0]}
            for source_row in (2, 3, 4, 5, 7, 8, 9):
                self.assertIsNone(by_original_row[source_row][4], f"source row {source_row}")
            self.assertEqual("01HAL11BP102", by_original_row[6][4])
            self.assertTrue(all((by_original_row[i + 2][1] or "") == row[0] for i, row in enumerate(rows)))
            self.assertIsNone(delivered["第二专业"].cell(2, 4).value)
            notes = [str(row[-1] or "") for row in by_original_row.values()]
            self.assertTrue(any("3.1.1" in note for note in notes))
            delivered.close()

    def test_physical_instruments_are_not_blank_and_quantity_rows_are_blank(self):
        for dtype in ("热电偶", "铠装壁温热电偶", "热电阻", "压力开关", "振动监测仪", "流量孔板"):
            with self.subTest(dtype=dtype):
                self.assertFalse(classify_row("01HAL11BT101", "出口温度1", dtype).is_point)
        self.assertTrue(classify_row("62LBG20FL001", "供汽母管流量", "测量设备").is_point)
        self.assertFalse(classify_row("62LBG20FL001", "供汽母管流量计", "测量设备").is_point)
        self.assertTrue(classify_row("01HAL11BT101", "温度反馈", "热电偶").is_point)

    def test_existing_final_code_never_gets_parent_from_conflicting_original(self):
        with tempfile.TemporaryDirectory() as directory:
            source = pathlib.Path(directory) / "parent.xlsx"
            folder = pathlib.Path(directory) / "out"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(HEADER + ["机组号", "原表公式"])
            sheet.append(["62LBG81QN001", "调门电源", "阀门执行机构", "61LBG81QN001", "调门电源", "", "1,2号机组公用"])
            sheet.append(["01HAL11BT101", "出口温度1", "热电偶", "", "", "", "1号机组", "=LEN(A3)"])
            sheet.append(["01HAL11BT102", "出口温度2", "热电阻", "", "", "", "1号机组"])
            workbook.save(source)
            workbook.close()
            with contextlib.redirect_stdout(io.StringIO()):
                gen_step1.main([str(source), str(folder), "--no-ai"])
            delivered = openpyxl.load_workbook(folder / "parent_已填写.xlsx")
            original = {row[0]: row for row in delivered.worksheets[0].iter_rows(min_row=2, values_only=True) if row[0]}
            self.assertEqual("61LBG81QN001", original[2][4])
            self.assertIsNone(original[2][6])
            self.assertIn("实际最终编码", original[2][-1])
            self.assertEqual("01HAL11BT101", original[3][4])
            self.assertEqual("01HAL11BT102", original[4][4])
            output_row = next(i for i, row in enumerate(delivered.worksheets[0].iter_rows(values_only=True), 1) if row[0] == 3)
            self.assertEqual(f"=LEN(B{output_row})", original[3][8])
            self.assertEqual("auto", delivered.calculation.calcMode)
            delivered.close()

    def test_preserved_bad_parent_blocks_generation_compliance_result(self):
        import json
        with tempfile.TemporaryDirectory() as directory:
            source = pathlib.Path(directory) / "bad-parent.xlsx"
            folder = pathlib.Path(directory) / "out"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(HEADER)
            sheet.append(["01HAL11GP001", "给水泵", "泵", "01HAL11GP001", "给水泵", "02HAL11"])
            workbook.save(source)
            workbook.close()
            with contextlib.redirect_stdout(io.StringIO()):
                gen_step1.main([str(source), str(folder), "--no-ai"])
            result = json.loads((folder / "第一步_生成结果.json").read_text(encoding="utf-8"))
            compliance = result["fill"]["compliance"]
            self.assertEqual("blocked", compliance["status"])
            self.assertTrue(any("父级" in error for issue in compliance["issues"] for error in issue["errors"]))

    def test_audit_detects_cross_sheet_duplicates_and_other_sheet_points(self):
        with tempfile.TemporaryDirectory() as directory:
            source = pathlib.Path(directory) / "cross.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(HEADER)
            sheet.append(["01HAL11GP001", "给水泵A", "泵", "01HAL11GP001", "给水泵A", "01HAL11"])
            other = workbook.create_sheet("第二专业")
            other.append(HEADER)
            other.append(["01HAL11GP001", "给水泵B", "泵", "01HAL11GP001", "给水泵B", "01HAL11"])
            other.append(["01HAL11GP002XB21", "运行反馈", "DI", "01HAL11GP002", "运行反馈", "01HAL11"])
            workbook.save(source)
            workbook.close()
            result = run_audit.audit_file(source, pathlib.Path(directory), requested_sheet="Sheet", run_ai_review=False, write_outputs=False)
            duplicates = [issue for issue in result["issues"] if issue["rule_id"] == "KKS-07"]
            self.assertEqual(2, len(duplicates))
            self.assertEqual({"Sheet", "第二专业"}, {issue["sheet"] for issue in duplicates})
            self.assertTrue(any(issue["rule_id"] == "KKS-38" and issue["sheet"] == "第二专业" for issue in result["issues"]))

    def test_audit_accepts_blank_points_and_reports_coded_points_without_mutation(self):
        with tempfile.TemporaryDirectory() as directory:
            source = pathlib.Path(directory) / "audit.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(HEADER)
            sheet.append(["01HAL11GP001XB21", "已运行反馈", "DI", "", "已运行反馈", "01HAL11"])
            sheet.append(["01HAL11BP101", "出口压力1", "AI", "", "出口压力反馈", "01HAL11"])
            sheet.append(["01HAL11GP001YB21", "启动指令", "DO", "01HAL11GP001", "启动指令", "01HAL11"])
            workbook.save(source)
            workbook.close()
            before = source.read_bytes()
            result = run_audit.audit_file(source, pathlib.Path(directory), run_ai_review=False, write_outputs=False)
            self.assertEqual(before, source.read_bytes())
            self.assertFalse(any(issue["rule_id"] == "KKS-00" and issue["excel_row"] in (2, 3)
                                 for issue in result["issues"]))
            violations = [issue for issue in result["issues"] if issue["rule_id"] == "KKS-38"]
            self.assertEqual([4], [issue["excel_row"] for issue in violations])
            self.assertTrue(violations[0]["compliance_locked"])
            self.assertEqual("blocked", result["compliance"]["status"])


if __name__ == "__main__":
    unittest.main()
