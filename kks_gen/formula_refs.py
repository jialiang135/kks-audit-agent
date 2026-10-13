"""Keep Excel formula references attached to original cells after insertion."""
import re
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import column_index_from_string, get_column_letter


def remap_formula(formula, row_map, column_shift, source_sheet, current_sheet):
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula
    tokens = Tokenizer(formula).items

    def endpoint(value):
        match = re.fullmatch(r"(\$?)([A-Za-z]{1,3})(\$?)([1-9][0-9]*)", value)
        if match:
            ca, column, ra, row = match.groups()
            return ca + get_column_letter(column_index_from_string(column) + column_shift) + ra + str(row_map.get(int(row), int(row)))
        match = re.fullmatch(r"(\$?)([A-Za-z]{1,3})", value)
        if match:
            return match[1] + get_column_letter(column_index_from_string(match[2]) + column_shift)
        match = re.fullmatch(r"(\$?)([1-9][0-9]*)", value)
        if match:
            return match[1] + str(row_map.get(int(match[2]), int(match[2])))
        return value

    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        value = token.value
        qualifier, separator, reference = value.rpartition("!")
        if separator:
            sheet = qualifier.strip("'").replace("''", "'")
            if sheet != source_sheet:
                continue
        else:
            reference = value
            if current_sheet != source_sheet:
                continue
        # Single bare names are named ranges, not entire-column references.
        if ":" not in reference and not re.fullmatch(r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]*", reference):
            continue
        token.value = (qualifier + "!" if separator else "") + ":".join(endpoint(part) for part in reference.split(":"))
    return "=" + "".join(token.value for token in tokens)
