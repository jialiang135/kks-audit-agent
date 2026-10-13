"""Device audit priority and removal of replacement-code generation."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl
import ai_review
import run_audit
import report_business
from kks_gen.audit_policy import findings, is_node


class AuditPolicyTest(unittest.TestCase):
    def record(self, code="01HAL11GP001", name="给水泵", dtype="泵", system="HAL"):
        return {"kks_code": code, "old_code": "", "name": name, "device_type": dtype, "system": system}

    def test_strict_device_format_and_nodes(self):
        for code in ("01HAL11GP001A", "K01HAL11GP001", "01HAL11GP01", "01HAL11GP001 "):
            checks = findings(self.record(code))
            self.assertEqual([0], [item[0] for item in checks])
            self.assertTrue(checks[0][3])
        self.assertFalse(is_node(self.record("01HAL11"), {"01HAL11"}))
        self.assertEqual([], findings(self.record("01HAL11", "给水系统", "系统")))

    def test_type_before_system_and_semantics(self):
        checks = findings(self.record("01HAL11GQ001", system="LBA"))
        self.assertEqual([1, 2, 3], [item[0] for item in checks])
        self.assertTrue(checks[0][3])
        self.assertTrue(checks[1][3])
        checks = findings(self.record(name="给水泵电机", dtype="电机"))
        self.assertEqual(1, checks[0][0])

    def test_valve_subject_after_pump_context(self):
        checks = findings(self.record("01HAL11QM001", name="给水泵出口电动门", dtype=""))
        self.assertFalse(any(stage == 1 and hard for stage, _, _, hard in checks))

    def test_differential_pressure_can_measure_flow(self):
        checks = findings(self.record("01HAL11BF101", name="给水流量", dtype="差压变送器"))
        self.assertFalse(any(stage == 1 and hard for stage, _, _, hard in checks))
        checks = findings(self.record("01HAL11BG001", name="磨煤机主电机与盘车连锁位", dtype=""))
        self.assertFalse(any(stage == 1 and hard for stage, _, _, hard in checks))

    def test_context_device_is_not_subject(self):
        for code, name in (("01HAL11QN001", "A密封风机进口电动调节门"),
                           ("01HAL11QM001", "A一次风机出口电动关断门"),
                           ("01HAL11BG001", "A一次风机动叶最小角"),
                           ("01HAL11EB201", "A一次风机润滑油站油箱A电加热器"),
                           ("01HAL11BL001", "A送风机轴承油箱油位低")):
            checks = findings(self.record(code, name=name, dtype=""))
            self.assertFalse(any(stage == 1 and hard for stage, _, _, hard in checks), name)

    def test_unknown_and_missing_evidence_stays_pending(self):
        checks = findings(self.record(name="装置", dtype="", system=""))
        self.assertTrue(all(not item[3] for item in checks))
        self.assertFalse(ai_review._is_candidate({"status": "needs_review", "evidence_missing": True}))

    def test_equipment_number_is_not_unit(self):
        from kks_gen.compliance import unit_conflicts
        self.assertEqual([], unit_conflicts("02LCD70BL104", "7号低压加热器液位4", "2号机组"))
        self.assertEqual([], unit_conflicts("01RDK21QA008", "SCR反应器B侧8号声波吹灰器", "1号机组"))
        self.assertTrue(unit_conflicts("01HAL11GP001", "2号机组给水泵"))

    def test_classification_reference_and_aggregate_context(self):
        ref = ai_review._classification_reference("01HAL11GP001")
        self.assertEqual("国标附录B", ref["device"]["source"])
        self.assertIn("流动", ref["device"]["definition"])
        result = {"_ai_context": {"records": [{"sheet": "一", "excel_row": 2, "kks_code": "", "name": "无关设备"}]}}
        packet = ai_review._build_evidence_packet(0, {"sheet": "一", "rule_id": "KKS-C"}, result)
        self.assertNotEqual("无关设备", packet["current_row"].get("name"))

    def test_multisheet_and_readonly_and_no_suggestion_report(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            book = openpyxl.Workbook()
            for sheet in [book.active, book.create_sheet("第二表")]:
                sheet.append(["KKS编码", "设备名称", "设备类型", "父级KKS编码", "系统"])
                sheet.append(["01HAL11GQ001", "给水泵", "泵", "01HAL11", "HAL"])
                sheet.append(["", "启动指令", "DO", "", ""])
            book.save(path)
            book.close()
            before = path.read_bytes()
            result = run_audit.audit_file(path, Path(directory), run_ai_review=False, write_outputs=False)
            self.assertEqual(before, path.read_bytes())
            self.assertEqual(2, len(result["audited_sheets"]))
            self.assertEqual(2, result["data_rows"])
            self.assertFalse(any(i["rule_id"] == "KKS-00" for i in result["issues"]))
            self.assertEqual({"Sheet", "第二表"}, {i["sheet"] for i in result["issues"] if i["rule_id"] == "KKS-40"})
            result["code_suggestions"] = {"status": "completed", "suggestions": {"2": {"suggested_code": "01HAL11GP003", "valid": True}}}
            overview = report_business.build_overview(result)
            self.assertNotIn("ai_code_suggestions", overview)
            html = report_business.render_html(result, overview)
            self.assertNotIn("AI建议编码", html)
            self.assertNotIn("01HAL11GP003", html)
            output = Path(directory) / "report.xlsx"
            report_business.render_xlsx(output, result, overview)
            exported = openpyxl.load_workbook(output)
            self.assertNotIn("AI建议编码", exported.sheetnames)
            exported.close()

    def test_ai_cannot_clear_strong_findings_or_missing_evidence(self):
        result = {"conclusion": "", "issues": [{"rule_id": "KKS-40", "status": "open", "compliance_locked": True},
                                                   {"rule_id": "KKS-45", "status": "needs_review", "evidence_missing": True}]}
        def fake_review(result, **kwargs):
            for item in result["issues"]:
                item.update(status="resolved", final_decision="likely_false_positive")
            return {"status": "completed"}
        with patch.object(run_audit, "review_issue_candidates", fake_review), patch.object(run_audit, "summarize_audit", return_value={}):
            run_audit.apply_ai_review_stage(result)
        self.assertEqual("confirmed_issue", result["issues"][0]["final_decision"])
        self.assertEqual("needs_human", result["issues"][1]["final_decision"])
        self.assertFalse(hasattr(ai_review, "suggest_correct_codes"))

    def test_evidence_uses_sheet_and_row(self):
        result = {"_ai_context": {"records": [{"sheet": "一", "excel_row": 2, "name": "泵"}, {"sheet": "二", "excel_row": 2, "name": "风机"}]}}
        packet = ai_review._build_evidence_packet(0, {"sheet": "二", "excel_row": 2}, result)
        self.assertEqual("风机", packet["current_row"]["name"])
