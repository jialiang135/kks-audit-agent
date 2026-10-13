"""Export original equipment rows whose delivered device code is blank."""
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from .row_policy import classify_row, is_table_instruction

PENDING_FILENAME = "设备编码待核清单.xlsx"


def write_pending_list(delivered_path, output_path, resolve_columns):
    source = openpyxl.load_workbook(delivered_path, read_only=True, data_only=False)
    result = openpyxl.Workbook()
    sheet = result.active
    sheet.title = "设备编码待核"
    sheet.append(["工作表", "源Excel行", "输出Excel行", "原KKS编码", "设备名称", "设备类型", "待核原因", "修编说明"])
    count = 0
    try:
        for ws in source.worksheets:
            iterator = iter(ws.iter_rows(values_only=True))
            prefix = []
            for _, row in zip(range(25), iterator):
                prefix.append(row)
            candidates = [(index, resolve_columns([str(value or "") for value in row]))
                          for index, row in enumerate(prefix)]
            candidates = [(index, columns) for index, columns in candidates if columns.get("kks") is not None]
            if not candidates:
                continue
            header_index, columns = max(candidates, key=lambda item: sum(value is not None for value in item[1].values()))
            header = prefix[header_index]
            seq_index = next((index for index, value in enumerate(header) if value == "原文件序号"), None)
            note_index = next((index for index, value in enumerate(header) if value == "修编说明"), None)
            from itertools import chain
            for output_row, row in enumerate(chain(prefix[header_index + 1:], iterator), header_index + 2):
                def get(key):
                    index = columns.get(key)
                    return row[index] if index is not None and index < len(row) else ""
                origin = row[seq_index] if seq_index is not None else output_row
                if not origin or get("kks") not in (None, ""):
                    continue
                old, name, dtype = get("design_code"), get("original_name") or get("device_name"), get("device_type")
                if not old and not name:
                    continue
                if is_table_instruction(get("kks"), name, old) or classify_row(old, name, dtype).is_point:
                    continue
                note = str(row[note_index] or "") if note_index is not None else ""
                reason = "重复编码待核" if "重复编码" in note or "撞成同一编码" in note else "分类/编号/工程约定待核"
                sheet.append([ws.title, origin, output_row, old, name, dtype, reason, note])
                # Source evidence is text; never execute a source formula here.
                for cell in sheet[count + 2]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                count += 1
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in zip("ABCDEFGH", (22, 14, 14, 25, 55, 24, 30, 100)):
            sheet.column_dimensions[column].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="254E70")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        result.save(output_path)
    finally:
        result.close()
        source.close()
    return {"file": Path(output_path).name, "rows": count}
